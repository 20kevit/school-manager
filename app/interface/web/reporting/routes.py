"""Reporting and Analytics Routes."""

from __future__ import annotations

import csv
import io
from datetime import date

import openpyxl
from flask import Blueprint, Response, render_template, request
from flask_login import current_user, login_required

from app.application.academy.class_service import list_classes
from app.application.academy.course_service import list_courses
from app.application.assessment.assessment_service import get_parent_children
from app.application.reporting.reporting_service import ReportingService
from app.interface.web.decorators import role_required

reporting_bp = Blueprint("reporting", __name__, url_prefix="/reporting")


@reporting_bp.route("/admin/overview")
@login_required
@role_required("ADMIN")
def admin_overview():
    """Admin dashboard overview with cross-system analytics."""
    overview = ReportingService.get_admin_overview()
    return render_template("reporting/admin_overview.html", overview=overview)


@reporting_bp.route("/admin/student/<int:student_id>")
@login_required
@role_required("ADMIN")
def admin_student_report(student_id: int):
    """Admin view of individual student report."""
    report = ReportingService.get_student_report(student_id)
    if "error" in report:
        return render_template("errors/404.html", message=report["error"]), 404
    return render_template(
        "reporting/admin_student_report.html", report=report, student_id=student_id
    )


@reporting_bp.route("/admin/class/<int:class_id>")
@login_required
@role_required("ADMIN")
def admin_class_report(class_id: int):
    """Admin view of class report."""
    report = ReportingService.get_class_report(class_id)
    if "error" in report:
        return render_template("errors/404.html", message=report["error"]), 404
    return render_template(
        "reporting/admin_class_report.html", report=report, class_id=class_id
    )


@reporting_bp.route("/admin/course/<int:course_id>")
@login_required
@role_required("ADMIN")
def admin_course_report(course_id: int):
    """Admin view of course report."""
    report = ReportingService.get_course_report(course_id)
    if "error" in report:
        return render_template("errors/404.html", message=report["error"]), 404
    return render_template(
        "reporting/admin_course_report.html", report=report, course_id=course_id
    )


@reporting_bp.route("/admin/attendance")
@login_required
@role_required("ADMIN")
def admin_attendance_report():
    """Admin attendance report with filters."""
    class_id = request.args.get("class_id", type=int)
    student_id = request.args.get("student_id", type=int)
    start_date_str = request.args.get("start_date")
    end_date_str = request.args.get("end_date")

    start_date = date.fromisoformat(start_date_str) if start_date_str else None
    end_date = date.fromisoformat(end_date_str) if end_date_str else None

    report = ReportingService.get_attendance_report(
        class_id=class_id,
        student_user_id=student_id,
        start_date=start_date,
        end_date=end_date,
    )

    classes = list_classes()
    alerts = ReportingService.get_attendance_alerts()

    return render_template(
        "reporting/admin_attendance_report.html",
        report=report,
        classes=classes,
        alerts=alerts,
        selected_class_id=class_id,
        selected_student_id=student_id,
        start_date=start_date_str,
        end_date=end_date_str,
    )


@reporting_bp.route("/admin/assessment/<int:assessment_id>")
@login_required
@role_required("ADMIN")
def admin_assessment_report(assessment_id: int):
    """Admin view of assessment report."""
    report = ReportingService.get_assessment_report(assessment_id)
    if "error" in report:
        return render_template("errors/404.html", message=report["error"]), 404
    return render_template(
        "reporting/admin_assessment_report.html",
        report=report,
        assessment_id=assessment_id,
    )


@reporting_bp.route("/admin/activity")
@login_required
@role_required("ADMIN")
def admin_activity_report():
    """Admin activity/homework report."""
    class_id = request.args.get("class_id", type=int)
    report = ReportingService.get_activity_report(class_id=class_id)
    classes = list_classes()

    return render_template(
        "reporting/admin_activity_report.html",
        report=report,
        classes=classes,
        selected_class_id=class_id,
    )


@reporting_bp.route("/admin/finance")
@login_required
@role_required("ADMIN")
def admin_finance_report():
    """Admin finance report with date filters."""
    start_date_str = request.args.get("start_date")
    end_date_str = request.args.get("end_date")

    start_date = date.fromisoformat(start_date_str) if start_date_str else None
    end_date = date.fromisoformat(end_date_str) if end_date_str else None

    report = ReportingService.get_finance_report(
        start_date=start_date, end_date=end_date
    )

    return render_template(
        "reporting/admin_finance_report.html",
        report=report,
        start_date=start_date_str,
        end_date=end_date_str,
    )


@reporting_bp.route("/admin/enrollment")
@login_required
@role_required("ADMIN")
def admin_enrollment_report():
    """Admin enrollment report with filters."""
    class_id = request.args.get("class_id", type=int)
    course_id = request.args.get("course_id", type=int)

    report = ReportingService.get_enrollment_report(
        class_id=class_id, course_id=course_id
    )
    classes = list_classes()
    courses = list_courses()

    return render_template(
        "reporting/admin_enrollment_report.html",
        report=report,
        classes=classes,
        courses=courses,
        selected_class_id=class_id,
        selected_course_id=course_id,
    )


@reporting_bp.route("/admin/teacher/<int:teacher_id>")
@login_required
@role_required("ADMIN")
def admin_teacher_report(teacher_id: int):
    """Admin view of teacher report."""
    report = ReportingService.get_teacher_report(teacher_id)
    if "error" in report:
        return render_template("errors/404.html", message=report["error"]), 404
    return render_template(
        "reporting/admin_teacher_report.html", report=report, teacher_id=teacher_id
    )


@reporting_bp.route("/admin/operational")
@login_required
@role_required("ADMIN")
def admin_operational_report():
    """Admin operational report."""
    report = ReportingService.get_operational_report()
    return render_template("reporting/admin_operational_report.html", report=report)


@reporting_bp.route("/admin/notification")
@login_required
@role_required("ADMIN")
def admin_notification_report():
    """Admin notification report."""
    report = ReportingService.get_notification_report()
    return render_template("reporting/admin_notification_report.html", report=report)


# Student routes
@reporting_bp.route("/student/overview")
@login_required
@role_required("STUDENT")
def student_overview():
    """Student personal report overview."""
    report = ReportingService.get_student_report(current_user.id)
    return render_template("reporting/student_overview.html", report=report)


# Parent routes
@reporting_bp.route("/parent/child/<int:child_id>")
@login_required
@role_required("PARENT")
def parent_child_report(child_id: int):
    """Parent view of child report."""
    children = get_parent_children(current_user.id)
    child_ids = [c.id for c in children]

    if child_id not in child_ids:
        return render_template("errors/403.html", message="Access denied"), 403

    report = ReportingService.get_student_report(child_id)
    if "error" in report:
        return render_template("errors/404.html", message=report["error"]), 404

    return render_template(
        "reporting/parent_child_report.html", report=report, child_id=child_id
    )


# Teacher routes
@reporting_bp.route("/teacher/overview")
@login_required
@role_required("TEACHER")
def teacher_overview():
    """Teacher personal report overview."""
    report = ReportingService.get_teacher_report(current_user.id)
    return render_template("reporting/teacher_overview.html", report=report)


@reporting_bp.route("/teacher/class/<int:class_id>")
@login_required
@role_required("TEACHER")
def teacher_class_report(class_id: int):
    """Teacher view of their class report."""
    from app.infrastructure.database.models import AcademyClassModel

    class_obj = AcademyClassModel.query.get(class_id)
    if not class_obj or class_obj.teacher_user_id != current_user.id:
        return render_template("errors/403.html", message="Access denied"), 403

    report = ReportingService.get_class_report(class_id)
    return render_template(
        "reporting/teacher_class_report.html", report=report, class_id=class_id
    )


# CSV Export routes
@reporting_bp.route("/admin/export/attendance")
@login_required
@role_required("ADMIN")
def export_attendance_csv():
    """Export attendance report as CSV."""
    class_id = request.args.get("class_id", type=int)
    student_id = request.args.get("student_id", type=int)

    report = ReportingService.get_attendance_report(
        class_id=class_id, student_user_id=student_id
    )

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Status", "Count"])
    writer.writerow(["Total", report["total"]])
    writer.writerow(["Present", report["present"]])
    writer.writerow(["Absent", report["absent"]])
    writer.writerow(["Late", report["late"]])
    writer.writerow(["Excused", report["excused"]])
    writer.writerow(["Percentage", f"{report['percentage']}%"])

    return Response(
        "\ufeff" + output.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=attendance_report.csv"},
    )


@reporting_bp.route("/admin/export/assessment/<int:assessment_id>")
@login_required
@role_required("ADMIN")
def export_assessment_csv(assessment_id: int):
    """Export assessment report as CSV."""
    report = ReportingService.get_assessment_report(assessment_id)
    if "error" in report:
        return render_template("errors/404.html", message=report["error"]), 404

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Metric", "Value"])
    writer.writerow(["Eligible Students", report["stats"]["eligible"]])
    writer.writerow(["Submitted", report["stats"]["submitted"]])
    writer.writerow(["Missing", report["stats"]["missing"]])
    writer.writerow(["Average", report["stats"]["average"]])
    writer.writerow(["Min", report["stats"]["min"]])
    writer.writerow(["Max", report["stats"]["max"]])
    writer.writerow(["Passed", report["stats"]["passed"]])
    writer.writerow(["Failed", report["stats"]["failed"]])
    writer.writerow(["Pass Rate", f"{report['stats']['pass_rate']}%"])

    output.write("\n\nScore Distribution\n")
    writer.writerow(["Range", "Count"])
    for range_name, count in report["stats"]["score_distribution"].items():
        writer.writerow([range_name, count])

    return Response(
        "\ufeff" + output.getvalue(),
        mimetype="text/csv",
        headers={
            "Content-Disposition": f"attachment; filename=assessment_{assessment_id}_report.csv"
        },
    )


@reporting_bp.route("/admin/export/finance")
@login_required
@role_required("ADMIN")
def export_finance_csv():
    """Export finance report as CSV."""
    start_date_str = request.args.get("start_date")
    end_date_str = request.args.get("end_date")

    start_date = date.fromisoformat(start_date_str) if start_date_str else None
    end_date = date.fromisoformat(end_date_str) if end_date_str else None

    report = ReportingService.get_finance_report(
        start_date=start_date, end_date=end_date
    )

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Metric", "Amount (Toman)"])
    writer.writerow(["Total Payments", report["total_payments"]])
    writer.writerow(["Total Charges", report["total_charges"]])
    writer.writerow(["Total Discounts", report["total_discounts"]])
    writer.writerow(["Total Refunds", report["total_refunds"]])
    writer.writerow(["Net Revenue", report["net_revenue"]])

    output.write("\n\nDebt by Student\n")
    writer.writerow(["Student ID", "Student Name", "Balance"])
    for item in report["debt_by_student"]:
        writer.writerow([item["student_id"], item["student_name"], item["balance"]])

    return Response(
        "\ufeff" + output.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=finance_report.csv"},
    )


# Excel Export routes
@reporting_bp.route("/admin/export/attendance/xlsx")
@login_required
@role_required("ADMIN")
def export_attendance_excel():
    """Export attendance report as Excel."""
    class_id = request.args.get("class_id", type=int)
    student_id = request.args.get("student_id", type=int)

    report = ReportingService.get_attendance_report(
        class_id=class_id, student_user_id=student_id
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance"
    ws.append(["Status", "Count"])
    ws.append(["Total", report["total"]])
    ws.append(["Present", report["present"]])
    ws.append(["Absent", report["absent"]])
    ws.append(["Late", report["late"]])
    ws.append(["Excused", report["excused"]])
    ws.append(["Percentage", f"{report['percentage']}%"])

    output = io.BytesIO()
    wb.save(output)
    wb.close()

    return Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=attendance_report.xlsx"},
    )


@reporting_bp.route("/admin/export/assessment/<int:assessment_id>/xlsx")
@login_required
@role_required("ADMIN")
def export_assessment_excel(assessment_id: int):
    """Export assessment report as Excel."""
    report = ReportingService.get_assessment_report(assessment_id)
    if "error" in report:
        return render_template("errors/404.html", message=report["error"]), 404

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Assessment"
    ws.append(["Metric", "Value"])
    ws.append(["Eligible Students", report["stats"]["eligible"]])
    ws.append(["Submitted", report["stats"]["submitted"]])
    ws.append(["Missing", report["stats"]["missing"]])
    ws.append(["Average", report["stats"]["average"]])
    ws.append(["Min", report["stats"]["min"]])
    ws.append(["Max", report["stats"]["max"]])
    ws.append(["Passed", report["stats"]["passed"]])
    ws.append(["Failed", report["stats"]["failed"]])
    ws.append(["Pass Rate", f"{report['stats']['pass_rate']}%"])

    ws.append([])
    ws.append(["Score Distribution"])
    ws.append(["Range", "Count"])
    for range_name, count in report["stats"]["score_distribution"].items():
        ws.append([range_name, count])

    output = io.BytesIO()
    wb.save(output)
    wb.close()

    return Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename=assessment_{assessment_id}_report.xlsx"
        },
    )


@reporting_bp.route("/admin/export/finance/xlsx")
@login_required
@role_required("ADMIN")
def export_finance_excel():
    """Export finance report as Excel."""
    start_date_str = request.args.get("start_date")
    end_date_str = request.args.get("end_date")

    start_date = date.fromisoformat(start_date_str) if start_date_str else None
    end_date = date.fromisoformat(end_date_str) if end_date_str else None

    report = ReportingService.get_finance_report(
        start_date=start_date, end_date=end_date
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Finance"
    ws.append(["Metric", "Amount (Toman)"])
    ws.append(["Total Payments", report["total_payments"]])
    ws.append(["Total Charges", report["total_charges"]])
    ws.append(["Total Discounts", report["total_discounts"]])
    ws.append(["Total Refunds", report["total_refunds"]])
    ws.append(["Net Revenue", report["net_revenue"]])

    ws.append([])
    ws.append(["Debt by Student"])
    ws.append(["Student ID", "Student Name", "Balance"])
    for item in report["debt_by_student"]:
        ws.append([item["student_id"], item["student_name"], item["balance"]])

    output = io.BytesIO()
    wb.save(output)
    wb.close()

    return Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=finance_report.xlsx"},
    )
