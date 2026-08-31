"""Tests for Phase 22: Excel Import/Export and Public Enrollment Portal."""

from __future__ import annotations

import io
from datetime import date

import openpyxl
import pytest

from app import create_app
from app.extensions import db as _db
from app.infrastructure.database.models import (
    AcademyClassModel,
    CourseModel,
    EnrollmentModel,
    RoleModel,
    StudentProfileModel,
    UserModel,
)


@pytest.fixture(scope="session")
def app():
    app = create_app("testing")
    return app


@pytest.fixture(scope="function")
def db(app):
    with app.app_context():
        _db.create_all()
        yield _db
        _db.session.rollback()
        _db.drop_all()


@pytest.fixture(scope="function")
def client(app, db):
    return app.test_client()


@pytest.fixture
def admin_user(db):
    user = UserModel(login="admin")
    user.set_password("adminpass123")
    role = RoleModel(name="ADMIN")
    db.session.add(role)
    db.session.flush()
    user.roles.append(role)
    db.session.add(user)
    db.session.commit()
    return user


@pytest.fixture
def student_user(db):
    user = UserModel(login="student1")
    user.set_password("studentpass123")
    role = RoleModel(name="STUDENT")
    db.session.add(role)
    db.session.flush()
    user.roles.append(role)
    db.session.add(user)
    db.session.flush()
    profile = StudentProfileModel(
        user_id=user.id,
        first_name="Ali",
        last_name="Rezaei",
        national_id="1234567890",
        mobile_phone="09121234567",
        birth_date=date(2000, 1, 15),
        gender="MALE",
    )
    db.session.add(profile)
    db.session.commit()
    return user


@pytest.fixture
def academy_class(db):
    course = CourseModel(
        title="Test Course",
        description="Test",
        num_sessions=10,
    )
    db.session.add(course)
    db.session.flush()
    cls = AcademyClassModel(
        course_id=course.id,
        title="Test Class",
        start_date=date(2026, 9, 1),
        end_date=date(2026, 12, 1),
        capacity=25,
        tuition=500000,
        num_sessions=10,
        status="ACTIVE",
    )
    db.session.add(cls)
    db.session.commit()
    return cls


def _make_excel(rows: list[list]) -> bytes:
    """Helper to create an Excel file from rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


# ===========================================================================
# Excel Parser Tests
# ===========================================================================


class TestExcelParser:
    """Tests for the Excel import parser."""

    def test_valid_import(self, db):
        from app.application.enrollment.excel_service import parse_excel_file

        data = _make_excel(
            [
                [
                    "first_name",
                    "last_name",
                    "national_id",
                    "mobile_phone",
                    "birth_date",
                    "gender",
                ],
                ["Reza", "Ahmadi", "1111111111", "09121111111", "1385/03/15", "MALE"],
                [
                    "Sara",
                    "Mohammadi",
                    "2222222222",
                    "09122222222",
                    "1386/05/20",
                    "FEMALE",
                ],
            ]
        )

        result = parse_excel_file(data)

        assert result.total_rows == 2
        assert result.created == 2
        assert result.errors == 0
        assert result.duplicates == 0

        users = UserModel.query.all()
        assert len(users) == 2
        profiles = StudentProfileModel.query.all()
        assert len(profiles) == 2

    def test_duplicate_national_id_skipped(self, db, student_user):
        from app.application.enrollment.excel_service import parse_excel_file

        data = _make_excel(
            [
                [
                    "first_name",
                    "last_name",
                    "national_id",
                    "mobile_phone",
                    "birth_date",
                ],
                ["Ali", "Rezaei", "1234567890", "09121234567", "1385/01/15"],
                ["New", "Student", "9999999999", "09129999999", "1386/05/20"],
            ]
        )

        result = parse_excel_file(data)

        assert result.total_rows == 2
        assert result.created == 1
        assert result.duplicates == 1

    def test_empty_file(self, db):
        from app.application.enrollment.excel_service import (
            ExcelImportError,
            parse_excel_file,
        )

        data = _make_excel([["first_name", "last_name"]])

        with pytest.raises(ExcelImportError, match="at least one data row"):
            parse_excel_file(data)

    def test_missing_required_columns(self, db):
        from app.application.enrollment.excel_service import (
            ExcelImportError,
            parse_excel_file,
        )

        data = _make_excel(
            [
                ["first_name", "last_name"],
                ["Ali", "Rezaei"],
            ]
        )

        with pytest.raises(ExcelImportError, match="Missing required columns"):
            parse_excel_file(data)

    def test_header_only_no_data(self, db):
        from app.application.enrollment.excel_service import (
            ExcelImportError,
            parse_excel_file,
        )

        data = _make_excel(
            [
                [
                    "first_name",
                    "last_name",
                    "national_id",
                    "mobile_phone",
                    "birth_date",
                ],
            ]
        )

        with pytest.raises(ExcelImportError, match="at least one data row"):
            parse_excel_file(data)

    def test_malformed_date_error(self, db):
        from app.application.enrollment.excel_service import parse_excel_file

        data = _make_excel(
            [
                [
                    "first_name",
                    "last_name",
                    "national_id",
                    "mobile_phone",
                    "birth_date",
                ],
                ["Ali", "Rezaei", "3333333333", "09123333333", "invalid-date"],
            ]
        )

        result = parse_excel_file(data)

        assert result.total_rows == 1
        assert result.errors == 1
        assert result.created == 0

    def test_missing_fields_error(self, db):
        from app.application.enrollment.excel_service import parse_excel_file

        data = _make_excel(
            [
                [
                    "first_name",
                    "last_name",
                    "national_id",
                    "mobile_phone",
                    "birth_date",
                ],
                ["", "Rezaei", "4444444444", "09124444444", "1385/03/15"],
            ]
        )

        result = parse_excel_file(data)

        assert result.total_rows == 1
        assert result.errors == 1
        assert result.created == 0

    def test_blank_rows_skipped(self, db):
        from app.application.enrollment.excel_service import parse_excel_file

        data = _make_excel(
            [
                [
                    "first_name",
                    "last_name",
                    "national_id",
                    "mobile_phone",
                    "birth_date",
                ],
                ["Reza", "Ahmadi", "5555555555", "09125555555", "1385/03/15"],
                [None, None, None, None, None],
                ["Sara", "Karimi", "6666666666", "09126666666", "1386/05/20"],
            ]
        )

        result = parse_excel_file(data)

        assert result.total_rows == 2
        assert result.created == 2

    def test_user_created_inactive_with_national_id_password(self, db):
        from app.application.enrollment.excel_service import parse_excel_file

        data = _make_excel(
            [
                [
                    "first_name",
                    "last_name",
                    "national_id",
                    "mobile_phone",
                    "birth_date",
                ],
                ["Reza", "Ahmadi", "7777777777", "09127777777", "1385/03/15"],
            ]
        )

        result = parse_excel_file(data)

        assert result.created == 1
        user = UserModel.query.filter_by(login="7777777777").first()
        assert user is not None
        assert user.has_role("STUDENT")
        assert user.check_password("7777777777")

    def test_gender_defaults_to_male(self, db):
        from app.application.enrollment.excel_service import parse_excel_file

        data = _make_excel(
            [
                [
                    "first_name",
                    "last_name",
                    "national_id",
                    "mobile_phone",
                    "birth_date",
                ],
                ["Reza", "Ahmadi", "8888888888", "09128888888", "1385/03/15"],
            ]
        )

        parse_excel_file(data)

        profile = StudentProfileModel.query.filter_by(national_id="8888888888").first()
        assert profile is not None
        assert profile.gender == "MALE"


# ===========================================================================
# Excel Export Tests
# ===========================================================================


class TestExcelExport:
    """Tests for Excel export functionality."""

    def test_export_students(self, db, student_user):
        from app.application.enrollment.excel_service import export_students_to_excel

        data = export_students_to_excel()
        assert isinstance(data, bytes)
        assert len(data) > 0

        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        assert len(rows) == 2  # header + 1 student
        assert rows[0][0] == "first_name"
        wb.close()


# ===========================================================================
# Public Enrollment Portal Tests
# ===========================================================================


class TestPublicEnrollment:
    """Tests for the public enrollment portal."""

    def test_public_apply_page_loads(self, client, db, academy_class):
        response = client.get("/enrollment/apply")
        assert response.status_code == 200
        assert "ثبت‌نام" in response.data.decode("utf-8")

    def test_public_apply_success(self, client, db, academy_class):
        response = client.post(
            "/enrollment/apply",
            data={
                "first_name": "Mohammad",
                "last_name": "Hosseini",
                "national_id": "9998887777",
                "mobile_phone": "09129998887",
                "birth_date": "1385/06/15",
                "gender": "MALE",
                "class_id": str(academy_class.id),
            },
            follow_redirects=True,
        )
        assert response.status_code == 200
        html = response.data.decode("utf-8")
        assert "ارسال شد" in html

        user = UserModel.query.filter_by(login="9998887777").first()
        assert user is not None
        assert not user.is_active_user
        assert user.has_role("STUDENT")

        profile = StudentProfileModel.query.filter_by(national_id="9998887777").first()
        assert profile is not None
        assert profile.first_name == "Mohammad"

        enrollment = EnrollmentModel.query.filter_by(student_user_id=user.id).first()
        assert enrollment is not None
        assert enrollment.status == "PENDING"

    def test_public_apply_existing_student(
        self, client, db, academy_class, student_user
    ):
        response = client.post(
            "/enrollment/apply",
            data={
                "first_name": "Ali",
                "last_name": "Rezaei",
                "national_id": "1234567890",
                "mobile_phone": "09121234567",
                "birth_date": "1378/01/15",
                "gender": "MALE",
                "class_id": str(academy_class.id),
            },
            follow_redirects=True,
        )
        assert response.status_code == 200

        enrollment = EnrollmentModel.query.filter_by(
            student_user_id=student_user.id
        ).first()
        assert enrollment is not None
        assert enrollment.status == "PENDING"

    def test_public_apply_missing_fields(self, client, db, academy_class):
        response = client.post(
            "/enrollment/apply",
            data={
                "first_name": "",
                "last_name": "Hosseini",
                "national_id": "9998887777",
                "mobile_phone": "09129998887",
                "birth_date": "1385/06/15",
                "gender": "MALE",
                "class_id": str(academy_class.id),
            },
            follow_redirects=True,
        )
        assert response.status_code == 200

        user = UserModel.query.filter_by(login="9998887777").first()
        assert user is None

    def test_public_apply_invalid_class(self, client, db):
        response = client.post(
            "/enrollment/apply",
            data={
                "first_name": "Mohammad",
                "last_name": "Hosseini",
                "national_id": "9998887777",
                "mobile_phone": "09129998887",
                "birth_date": "1385/06/15",
                "gender": "MALE",
                "class_id": "9999",
            },
            follow_redirects=True,
        )
        assert response.status_code == 200

    def test_public_apply_duplicate_enrollment(
        self, client, db, academy_class, student_user
    ):
        client.post(
            "/enrollment/apply",
            data={
                "first_name": "Ali",
                "last_name": "Rezaei",
                "national_id": "1234567890",
                "mobile_phone": "09121234567",
                "birth_date": "1378/01/15",
                "gender": "MALE",
                "class_id": str(academy_class.id),
            },
            follow_redirects=True,
        )

        response = client.post(
            "/enrollment/apply",
            data={
                "first_name": "Ali",
                "last_name": "Rezaei",
                "national_id": "1234567890",
                "mobile_phone": "09121234567",
                "birth_date": "1378/01/15",
                "gender": "MALE",
                "class_id": str(academy_class.id),
            },
            follow_redirects=True,
        )
        assert response.status_code == 200

    def test_approve_activate_enrollment(self, client, db, admin_user, academy_class):
        client.post(
            "/enrollment/apply",
            data={
                "first_name": "Mohammad",
                "last_name": "Hosseini",
                "national_id": "9998887777",
                "mobile_phone": "09129998887",
                "birth_date": "1385/06/15",
                "gender": "MALE",
                "class_id": str(academy_class.id),
            },
        )

        user = UserModel.query.filter_by(login="9998887777").first()
        enrollment = EnrollmentModel.query.filter_by(student_user_id=user.id).first()

        client.post("/auth/login", data={"login": "admin", "password": "adminpass123"})
        response = client.post(
            f"/enrollment/admin/{enrollment.id}/approve-activate",
            follow_redirects=True,
        )
        assert response.status_code == 200

        user = UserModel.query.get(user.id)
        assert user.is_active_user

        enrollment = EnrollmentModel.query.get(enrollment.id)
        assert enrollment.status == "APPROVED"

    def test_admin_import_page_loads(self, client, db, admin_user):
        client.post("/auth/login", data={"login": "admin", "password": "adminpass123"})
        response = client.get("/enrollment/admin/import/students")
        assert response.status_code == 200
        assert "ورود اطلاعات" in response.data.decode("utf-8")

    def test_admin_import_excel_upload(self, client, db, admin_user):
        data = _make_excel(
            [
                [
                    "first_name",
                    "last_name",
                    "national_id",
                    "mobile_phone",
                    "birth_date",
                ],
                ["Reza", "Ahmadi", "1111111111", "09121111111", "1385/03/15"],
                ["Sara", "Mohammadi", "2222222222", "09122222222", "1386/05/20"],
            ]
        )

        client.post("/auth/login", data={"login": "admin", "password": "adminpass123"})
        response = client.post(
            "/enrollment/admin/import/students",
            data={"file": (io.BytesIO(data), "students.xlsx")},
            content_type="multipart/form-data",
            follow_redirects=True,
        )
        assert response.status_code == 200

        profiles = StudentProfileModel.query.all()
        assert len(profiles) == 2

    def test_admin_import_no_file(self, client, db, admin_user):
        client.post("/auth/login", data={"login": "admin", "password": "adminpass123"})
        response = client.post(
            "/enrollment/admin/import/students",
            data={},
            content_type="multipart/form-data",
            follow_redirects=True,
        )
        assert response.status_code == 200

    def test_pending_list_shows_public_enrollments(
        self, client, db, admin_user, academy_class
    ):
        client.post(
            "/enrollment/apply",
            data={
                "first_name": "Mohammad",
                "last_name": "Hosseini",
                "national_id": "9998887777",
                "mobile_phone": "09129998887",
                "birth_date": "1385/06/15",
                "gender": "MALE",
                "class_id": str(academy_class.id),
            },
        )

        client.post("/auth/login", data={"login": "admin", "password": "adminpass123"})
        response = client.get("/enrollment/admin/pending")
        assert response.status_code == 200
        html = response.data.decode("utf-8")
        assert "Hosseini" in html or "Hossein" in html or "9998887777" in html
